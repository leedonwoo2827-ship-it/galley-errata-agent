"""엑셀 정오표.

unit_book(): 단위 정오표 — [정오표] 반영 건 · [전체 후보] 체크하지 않은 것까지
merge():     workspace/정오표_<책>_전체.xlsx + 같은 이름 .html (한 번에 보는 책 전체 정오표)
  [정오표]      검토완료 단위는 반영 건, 검토 전 단위는 기본 체크 기준으로 자동 선정한 건
  [검토 후보]   선정되지 않은 나머지 후보 — 엑셀로 나눠 정독할 때 쓴다
  [진행 현황]   단위 · 쪽 · 담당자 · 마감 · 상태 · 건수
  [책 전체 메모] 규칙 검사의 책 전체 메모
"""
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .report import checked

HEAD_FILL = PatternFill("solid", fgColor="E9EEF6")
THIN = Side(style="thin", color="C9D1DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
RED, GREEN = Font(color="B42318"), Font(color="05603A", bold=True)

ERRATA_COLS = [("단위", 8), ("쪽", 6), ("유형", 10), ("기존", 45), ("수정", 45), ("사유", 40),
               ("출처", 6), ("수정확률", 8), ("상태", 16), ("반영자", 9), ("메모", 24)]


def _sheet(ws, cols, rows, red=None, green=None):
    ws.append([c for c, _ in cols])
    for k, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
        c = ws.cell(row=1, column=k)
        c.font, c.fill, c.border, c.alignment = Font(bold=True), HEAD_FILL, BOX, Alignment(vertical="center")
    for r in rows:
        ws.append(r)
        i = ws.max_row
        for k in range(1, len(cols) + 1):
            c = ws.cell(row=i, column=k)
            c.alignment, c.border = WRAP, BOX
        if red:
            ws.cell(row=i, column=red).font = RED
        if green:
            ws.cell(row=i, column=green).font = GREEN
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def _row(r, reviewer=""):
    return [r.get("unit", ""), r["page"], r["judge_type"], r["old_text"],
            r["new_text"] if r["new_text"].strip() else "(삭제)", r["reason"],
            "규칙" if r["source"] == "rule" else "AI", round(r["p_revise"] * 100),
            r.get("status", "검토 전"), reviewer, r.get("note", "")]


def is_errata(r):
    return r.get("apply") and r.get("status") != "변경없음"


def unit_book(path, records, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "정오표"
    _sheet(ws, ERRATA_COLS, [_row(r, meta["reviewer"]) for r in records if is_errata(r)], red=4, green=5)
    _sheet(wb.create_sheet("전체 후보"), ERRATA_COLS,
           [_row(r, meta["reviewer"] if r.get("apply") else "") for r in records], red=4, green=5)
    info = wb.create_sheet("정보")
    for k, v in (("단위", meta["unit"]), ("제목", meta["title"]), ("인쇄 쪽", f"{meta['printed'][0]}–{meta['printed'][1]}"),
                 ("반영자", meta["reviewer"]), ("작성", meta["date"]), ("담당자", meta.get("assignee", "")),
                 ("마감", meta.get("due", ""))):
        info.append([k, v])
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 10, 60
    wb.save(path)
    return path


def unit_state(u):
    """단위 진행 상태 + (반영 기록 or 판정 후보)."""
    d = config.unit_dir(u["id"])
    latest = d / "05_정오표" / "latest.json"
    res = d / "03_판정" / "results.json"
    if latest.exists():
        lt = json.loads(latest.read_text(encoding="utf-8"))
        return "검토완료", lt["records"], lt
    if res.exists():
        # 검토 전 단위는 검토 화면의 기본 체크 기준으로 자동 선정해 둔다 (한 번에 전체 정오표가 나오게)
        rows = json.loads(res.read_text(encoding="utf-8"))["rows"]
        return "판정완료", [{**r, "unit": u["id"], "apply": checked(r),
                             "status": "자동 선정(검토 전)" if checked(r) else "검토 전"} for r in rows], None
    return "미판정", [], None


def merge():
    units = [u for u in config.units() if u.get("include", True)]
    errata, cands, prog = [], [], []
    for u in units:
        state, recs, lt = unit_state(u)
        reviewer = (lt or {}).get("reviewer", "")
        errata += [_row(r, reviewer) for r in recs if is_errata(r)]
        cands += [_row(r) for r in recs if not is_errata(r) and r.get("status") != "변경없음"]
        prog.append([u["id"], u["title"], f"{u['printed'][0]}–{u['printed'][1]}", u.get("n_pages", ""),
                     u.get("assignee", ""), u.get("due", ""), state, len(recs),
                     sum(1 for r in recs if is_errata(r)), reviewer, (lt or {}).get("date", "")])
    wb = Workbook()
    ws = wb.active
    ws.title = "정오표"
    _sheet(ws, ERRATA_COLS, errata, red=4, green=5)
    _sheet(wb.create_sheet("검토 후보"), ERRATA_COLS, cands, red=4, green=5)
    _sheet(wb.create_sheet("진행 현황"),
           [("단위", 8), ("제목", 48), ("인쇄 쪽", 10), ("쪽수", 6), ("담당자", 10), ("마감", 11), ("상태", 10),
            ("후보", 6), ("정오표", 7), ("반영자", 9), ("반영일시", 16)], prog)
    notes = wb.create_sheet("책 전체 메모")
    notes.append(["메모"]); notes["A1"].font = Font(bold=True)
    for n in (config.rules_book() or {}).get("notes", []):
        notes.append([n])
    notes.column_dimensions["A"].width = 140
    notes.append([])
    notes.append([f"작성 {datetime.now():%Y-%m-%d %H:%M} · {(config.toc() or {}).get('book', '')}"])
    path = config.merged_xlsx()
    wb.save(path)
    board_html(errata, prog)
    print(f"[엑셀] 전체 정오표 — {len(errata)}건 (검토완료 {sum(p[6] == '검토완료' for p in prog)}개 단위 · "
          f"나머지는 자동 선정) · 검토 후보 {len(cands)}건 → {path}")
    return path


def board_html(errata, prog):
    """책 한 권 정오표를 한 장의 HTML로 (게시판·가제본 배포용). 단위 순서 → 쪽 순서."""
    import html
    from .apply import page_html
    e = html.escape
    trs = "".join(
        f"<tr><td>{e(str(r[0]))}</td><td class=n>{r[1]}</td><td>{e(r[2])}</td><td><del>{e(r[3])}</del></td>"
        f"<td>{'<i>(삭제)</i>' if r[4] == '(삭제)' else '<ins>' + e(r[4]) + '</ins>'}</td><td>{e(r[5])}</td>"
        f"<td class=s>{e(r[8])}</td></tr>" for r in errata)
    done = sum(p[6] == "검토완료" for p in prog)
    body = (f"<h1>정오표 — {e((config.toc() or {}).get('book', ''))}</h1>"
            f"<p class=meta>작성 {datetime.now():%Y-%m-%d %H:%M} · {len(errata)}건 · 검토완료 {done}/{len(prog)}개 단위"
            f"<br>상태가 '자동 선정(검토 전)'인 줄은 아직 사람이 확인하지 않은 AI·규칙 판정입니다.</p>"
            f"<table><thead><tr><th>단위</th><th>쪽</th><th>유형</th><th>기존</th><th>수정</th><th>사유</th><th>상태</th></tr></thead>"
            f"<tbody>{trs or '<tr><td colspan=7>아직 정오표 항목이 없습니다.</td></tr>'}</tbody></table>")
    out = config.merged_xlsx().with_suffix(".html")
    out.write_text(page_html("정오표 전체", body).replace("</style>", "td.s{font-size:12px;color:#666;white-space:nowrap}</style>"),
                   encoding="utf-8")
    return out


if __name__ == "__main__":
    merge()
